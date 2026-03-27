import markdown
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
import re
import os

def convert_md_to_docx(md_file_path, docx_file_path):
    # Read Markdown content
    with open(md_file_path, 'r', encoding='utf-8') as f:
        md_content = f.read()

    # Create a new Document
    doc = Document()
    
    # Set default font
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(12)
    
    # Configure paragraph spacing for the Normal style
    paragraph_format = style.paragraph_format
    paragraph_format.space_after = Pt(12)  # Add space after paragraphs
    paragraph_format.line_spacing = 1.15   # Slightly increased line spacing

    # Process content line by line to handle basic formatting manually
    # This is a simple parser to handle the specific structure of the chapter
    lines = md_content.split('\n')
    
    for line in lines:
        line = line.strip()
        if not line:
            continue
            
        # Handle Headers
        if line.startswith('# '):
            p = doc.add_heading(line[2:], level=1)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.runs[0]
            run.font.name = 'Times New Roman'
            run.font.color.rgb = None # Reset color to default
            continue
            
        if line.startswith('## '):
            p = doc.add_heading(line[3:], level=2)
            run = p.runs[0]
            run.font.name = 'Times New Roman'
            run.font.color.rgb = None
            continue
            
        # Handle Images (skip them but add a placeholder note or just skip)
        if line.startswith('![') and '](' in line:
            # Extract alt text
            alt_text = line[2:line.find(']')]
            p = doc.add_paragraph(f"[Image: {alt_text}]")
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.runs[0].italic = True
            continue
            
        # Handle Separators
        if line == '***' or line == '---':
            p = doc.add_paragraph('***')
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            continue
            
        # Handle Bold Text (simple implementation)
        if line.startswith('**') and line.endswith('**'):
            p = doc.add_paragraph()
            run = p.add_run(line[2:-2])
            run.bold = True
            # Check if it's a location/date line or a section title
            if "Glassworthy Manor" in line or "The Green Equation" in line or "The Logic of the Lute" in line or "The Arrival of the News" in line:
                 p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            continue

        # Regular Paragraphs
        p = doc.add_paragraph()
        
        # Simple inline formatting parser
        # This handles *italics* and **bold** within the text
        # Note: This is a basic parser and might not handle nested or complex cases perfectly
        
        parts = re.split(r'(\*\*.*?\*\*|\*.*?\*)', line)
        for part in parts:
            if part.startswith('**') and part.endswith('**'):
                run = p.add_run(part[2:-2])
                run.bold = True
            elif part.startswith('*') and part.endswith('*'):
                run = p.add_run(part[1:-1])
                run.italic = True
            else:
                p.add_run(part)
                
        # Add extra spacing for specific paragraphs if needed, 
        # though the style default handles most cases.

    # Save the document
    doc.save(docx_file_path)
    print(f"Successfully converted {md_file_path} to {docx_file_path}")

if __name__ == "__main__":
    md_path = "/home/ubuntu/durban-paradox/client/public/content/chapters/1.md"
    docx_path = "/home/ubuntu/durban-paradox/client/public/content/chapters/Chapter_1_The_Lonely_Number.docx"
    convert_md_to_docx(md_path, docx_path)
